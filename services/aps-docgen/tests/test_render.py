"""Tests for the merge, built around the ways this can silently go wrong.

Test-only. Nothing here is imported by `app.py` or `render.py`; pytest collects
the `test_*` functions and that is the whole of its role.

Two deliberate choices about fixtures:

- Every .docx is generated in-process from plain strings, so a failure points
  at a readable cause rather than an opaque committed binary.
- `JOB_TOKENS` is FICTIONAL data in the SHAPE of the real payload. The shape
  matters — a parent that is both a scalar and a namespace is exactly what
  breaks — but the identities must not be real client details in a repo.
"""

from __future__ import annotations

import io
import zipfile

import pytest
from docx import Document

from context import Value, build_context
from render import RenderError, render_docx


def make_docx(paragraphs: list[str]) -> bytes:
    """A .docx whose paragraphs are exactly the given strings."""
    doc = Document()
    for text in paragraphs:
        doc.add_paragraph(text)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def make_docx_split_runs(fragments: list[str]) -> bytes:
    """A .docx with ONE paragraph whose text is split across separate runs.

    This is what Word actually produces: spellcheck, autocorrect and revision
    ids fragment typed text, so ``{{ Client }}`` can end up as ``{{ Cli`` +
    ``ent }}`` in the XML. If the merge cannot cope, tokens fail seemingly at
    random — the failure mode looks like a bug in our code, not the template.
    """
    doc = Document()
    paragraph = doc.add_paragraph()
    for fragment in fragments:
        paragraph.add_run(fragment)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def make_docx_table(rows: list[list[str]]) -> bytes:
    doc = Document()
    table = doc.add_table(rows=len(rows), cols=max(len(r) for r in rows))
    for r, row in enumerate(rows):
        for c, text in enumerate(row):
            table.cell(r, c).text = text
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def text_of(docx_bytes: bytes) -> str:
    doc = Document(io.BytesIO(docx_bytes))
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            parts.extend(cell.text for cell in row.cells)
    return "\n".join(parts)


# Fictional, but shaped like the live payload: `Client` is both a scalar and
# the parent of `Client.Mail`, which is the case that breaks naive designs.
JOB_TOKENS = [
    {"name": "Title", "value": "0000001"},
    {"name": "Client", "value": "Example Statutory Authority"},
    {"name": "Client.Mail", "value": "notices@example.test"},
    {"name": "Client.Telephone", "value": "02 0000 0000"},
    {"name": "Parties", "value": "Jane Example"},
    {"name": "Hearing_Date", "value": "25 July 2026"},
]


# --------------------------------------------------------------------------
# The collision: a name that is both a scalar and a namespace.
# --------------------------------------------------------------------------


def test_parent_renders_as_scalar_and_resolves_children():
    """The single most important behaviour in this service.

    All 14 parents in the live data are both, so if this breaks, every
    template that mentions a client or a party breaks with it.
    """
    template = make_docx([
        "Client: {{ Client }}",
        "Email: {{ Client.Mail }}",
        "Phone: {{ Client.Telephone }}",
    ])
    out = text_of(render_docx(template, JOB_TOKENS))

    assert "Client: Example Statutory Authority" in out
    assert "Email: notices@example.test" in out
    assert "Phone: 02 0000 0000" in out


def test_parent_scalar_is_truthy_for_if():
    """`{% if Client %}` must follow the string, not the row count.

    This is why `Value.__len__` is deliberately not defined — defining it
    would make a populated client with no table rows read as empty.
    """
    template = make_docx([
        "{% if Client %}HAS CLIENT{% endif %}",
        "{% if Missing_Field %}NO{% else %}EMPTY OK{% endif %}",
    ])
    out = text_of(render_docx(template, JOB_TOKENS))

    assert "HAS CLIENT" in out
    assert "EMPTY OK" in out


def test_string_filters_still_work_on_values():
    """Values must behave as strings — authors will reach for filters."""
    template = make_docx(["{{ Client.Mail|upper }}", "{{ Title|trim }}"])
    out = text_of(render_docx(template, JOB_TOKENS))

    assert "NOTICES@EXAMPLE.TEST" in out
    assert "0000001" in out


# --------------------------------------------------------------------------
# Authoring mistakes must degrade, not explode.
# --------------------------------------------------------------------------


def test_unknown_token_renders_empty_without_failing():
    template = make_docx(["Start[{{ Not_A_Real_Token }}]End"])
    out = text_of(render_docx(template, JOB_TOKENS))

    assert "Start[]End" in out


def test_unknown_nested_token_renders_empty_without_failing():
    """ChainableUndefined earns its place here.

    With Jinja's default Undefined this raises and the whole affidavit fails
    over one mistyped token.
    """
    template = make_docx(["Start[{{ Nope.Also_Nope }}]End"])
    out = text_of(render_docx(template, JOB_TOKENS))

    assert "Start[]End" in out


def test_malformed_template_raises_clean_error():
    template = make_docx(["{% for x in Client %}unclosed"])
    with pytest.raises(RenderError) as excinfo:
        render_docx(template, JOB_TOKENS)

    assert "placeholder is malformed" in str(excinfo.value)


def test_non_docx_input_raises_clean_error():
    """A corrupt file must not be reported as a template authoring mistake.

    docxtpl defers the file load, so without an explicit container check this
    surfaces during render() and tells the author to fix `{% for %}` tags in a
    template that is perfectly fine.
    """
    with pytest.raises(RenderError) as excinfo:
        render_docx(b"this is not a docx", JOB_TOKENS)

    message = str(excinfo.value)
    assert "not a Word .docx file" in message
    assert "placeholder" not in message


def test_zip_without_document_part_raises_clean_error():
    """A .doc or .rtf renamed to .docx is the realistic version of this.

    Some are valid zips, so the container check has to look for the OOXML
    document part, not just successful unzipping.
    """
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as archive:
        archive.writestr("something/else.xml", "<x/>")

    with pytest.raises(RenderError) as excinfo:
        render_docx(buf.getvalue(), JOB_TOKENS)

    assert "no word/document.xml part" in str(excinfo.value)


# --------------------------------------------------------------------------
# Word's run splitting — the empirical question.
# --------------------------------------------------------------------------


def test_token_split_across_runs_still_merges():
    """Does docxtpl's own patch_xml handle fragmented runs, or must we?

    The plan called a custom run-merge pass "mandatory". This test decides it
    on evidence: if it passes, the pass is redundant and we skip the
    complexity; if it fails, we write one.
    """
    template = make_docx_split_runs(["Client: {{ Cli", "ent", " }}"])
    out = text_of(render_docx(template, JOB_TOKENS))

    assert "Client: Example Statutory Authority" in out


def test_nested_token_split_across_runs_still_merges():
    template = make_docx_split_runs(["Email: {{ Client.", "Ma", "il }}"])
    out = text_of(render_docx(template, JOB_TOKENS))

    assert "Email: notices@example.test" in out


# --------------------------------------------------------------------------
# Dynamic table rows — in scope per the 2026-08-26 decision.
# --------------------------------------------------------------------------


def test_dynamic_table_rows_expand():
    """`{%tr for %}` repeats the table row per fee.

    The flat token list cannot express this — the Drupal resolver flattens
    multi-value fields to "Delivery Fee, Collection Fee" — so rows arrive in
    the separate `tables` payload.
    """
    template = make_docx_table([
        ["{%tr for fee in Job_Fees %}", ""],
        ["{{ fee.Description }}", "{{ fee.Amount }}"],
        ["{%tr endfor %}", ""],
    ])
    tables = {
        "Job_Fees": [
            {"Description": "Delivery Fee", "Amount": "35.00"},
            {"Description": "Collection Fee", "Amount": "34.94"},
        ]
    }
    out = text_of(render_docx(template, JOB_TOKENS, tables))

    assert "Delivery Fee" in out
    assert "35.00" in out
    assert "Collection Fee" in out
    assert "34.94" in out
    assert "{%tr" not in out


def test_table_parent_is_still_a_scalar_and_countable():
    template = make_docx([
        "Fees summary: {{ Job_Fees }}",
        "Count: {{ Job_Fees.rows|length }}",
    ])
    tokens = JOB_TOKENS + [{"name": "Job_Fees", "value": "Delivery Fee, Collection Fee"}]
    tables = {"Job_Fees": [{"Description": "Delivery Fee"}, {"Description": "Collection Fee"}]}
    out = text_of(render_docx(template, tokens, tables))

    assert "Fees summary: Delivery Fee, Collection Fee" in out
    assert "Count: 2" in out


def test_loop_over_scalar_yields_nothing_not_characters():
    """A template that loops a scalar is an authoring error.

    Yielding nothing drops the row; str's default iteration would explode the
    table into one row per letter.
    """
    template = make_docx(["{% for x in Client %}[{{ x }}]{% endfor %}DONE"])
    out = text_of(render_docx(template, JOB_TOKENS))

    assert "DONE" in out
    assert "[E]" not in out


# --------------------------------------------------------------------------
# Context construction.
# --------------------------------------------------------------------------


def test_shadowing_guard_rejects_lowercase_child_names():
    """The guard that keeps `__getattr__` safe.

    A child called `title` would be shadowed by `str.title` and render as a
    method repr, so it must fail loudly at build time.
    """
    with pytest.raises(ValueError, match="shadowed"):
        build_context([{"name": "Client.title", "value": "x"}])


def test_missing_parent_scalar_still_allows_children():
    """Drupal may emit `Client.Mail` with no `Client` pair at all."""
    context = build_context([{"name": "Client.Mail", "value": "a@example.test"}])

    assert context["Client"] == ""
    assert context["Client"].Mail == "a@example.test"


def test_values_coerce_to_document_text():
    context = build_context([
        {"name": "Empty", "value": None},
        {"name": "Number", "value": 42},
        {"name": "Flag", "value": True},
        {"name": "Many", "value": ["a", "b"]},
    ])

    assert context["Empty"] == ""
    assert context["Number"] == "42"
    assert context["Flag"] == "Yes"
    assert context["Many"] == "a, b"


def test_value_is_a_real_string():
    value = Value("hello", children={"Sub": Value("world")})

    assert isinstance(value, str)
    assert value.upper() == "HELLO"
    assert value.Sub == "world"
    assert value.Missing == ""


# --------------------------------------------------------------------------
# One-column lists: one paragraph per item from a plain {{ }} token.
# --------------------------------------------------------------------------


def make_docx_styled(paragraphs: list[tuple[str, str]]) -> bytes:
    """A .docx whose paragraphs carry a named Word style (e.g. List Number)."""
    doc = Document()
    for text, style in paragraphs:
        doc.add_paragraph(text, style=style)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def paragraphs_of(docx_bytes: bytes) -> list[str]:
    return [p.text for p in Document(io.BytesIO(docx_bytes)).paragraphs]


# The two documents on a Job, as `resolveListTokensForJob()` sends them. This
# is the case the feature exists for: the flat pair for this name arrives
# EMPTY, because Drupal's legacy resolver renders the parent field once and
# gives that render to every child of it.
DOCUMENT_LISTS = {
    "Documents.Affidavit_Template": [
        "Example Affidavit of Service",
        "Example Application for Minor Civil Dispute",
    ]
}


def test_list_token_renders_one_paragraph_per_item():
    """The whole point: a plain {{ }} token, one document per row.

    The author writes no loop and builds no table — this has to work from the
    same token the picker already offers, because templates are authored in
    Word for the web by staff, not by developers.
    """
    template = make_docx(["{{ Documents.Affidavit_Template }}"])
    out = paragraphs_of(render_docx(template, JOB_TOKENS, None, DOCUMENT_LISTS))

    assert "Example Affidavit of Service" in out
    assert "Example Application for Minor Civil Dispute" in out
    # Separate paragraphs, not one line with both names run together.
    assert out.index("Example Application for Minor Civil Dispute") == (
        out.index("Example Affidavit of Service") + 1
    )


def test_list_token_overrides_the_flat_pair():
    """Deliberate precedence, and the reason the feature is not additive.

    The flat pair for this name is the broken value — every document's whole
    rendered line glued together with no separator. The list must win.
    """
    tokens = JOB_TOKENS + [
        {
            "name": "Documents.Affidavit_Template",
            "value": "Example Affidavit of Serviceissued onExample Application",
        }
    ]
    template = make_docx(["{{ Documents.Affidavit_Template }}"])
    out = paragraphs_of(render_docx(template, tokens, None, DOCUMENT_LISTS))

    assert "Example Affidavit of Serviceissued onExample Application" not in out
    assert "Example Affidavit of Service" in out


def test_single_item_list_makes_one_paragraph_not_a_blank_line():
    """A Job with one document must not gain a trailing empty bullet."""
    template = make_docx(["{{ Documents.Affidavit_Template }}"])
    lists = {"Documents.Affidavit_Template": ["Only Template"]}
    out = paragraphs_of(render_docx(template, JOB_TOKENS, None, lists))

    assert out.count("Only Template") == 1
    assert out.count("") == 0


def test_empty_list_does_not_blank_an_existing_value():
    """An absent list is silence, not an instruction to clear the token."""
    tokens = JOB_TOKENS + [{"name": "Documents.Affidavit_Template", "value": "Fallback"}]
    template = make_docx(["{{ Documents.Affidavit_Template }}"])
    out = text_of(render_docx(template, tokens, None, {"Documents.Affidavit_Template": []}))

    assert "Fallback" in out


def test_blank_items_are_dropped_not_rendered_as_empty_rows():
    """A document with no template must leave no empty bullet behind."""
    template = make_docx(["{{ Documents.Affidavit_Template }}"])
    lists = {"Documents.Affidavit_Template": ["First", "", None, "   ", "Second"]}
    out = paragraphs_of(render_docx(template, JOB_TOKENS, None, lists))

    assert out == ["First", "Second"]


def test_list_items_keep_the_paragraph_style_so_numbering_works():
    """Each item must inherit the <w:pPr> of the paragraph the token sat in.

    This is why the separator is "\\a" and not "\\n": a soft line break stays
    inside one paragraph, so a numbered list would show a single "1." with
    every document under it. Real paragraphs each get their own number.
    """
    template = make_docx_styled([("{{ Documents.Affidavit_Template }}", "List Number")])
    doc = Document(io.BytesIO(render_docx(template, JOB_TOKENS, None, DOCUMENT_LISTS)))
    styled = [p for p in doc.paragraphs if p.text]

    assert len(styled) == 2
    assert all(p.style.name == "List Number" for p in styled)


def test_list_token_still_loops_if_the_author_wrote_one():
    """`lists` must not disturb the `{%tr for %}` path that already worked."""
    template = make_docx_table([
        ["{%tr for fee in Job_Fees %}", ""],
        ["{{ fee.Description }}", "{{ fee.Amount }}"],
        ["{%tr endfor %}", ""],
    ])
    tables = {"Job_Fees": [{"Description": "Delivery Fee", "Amount": "35.00"}]}
    out = text_of(render_docx(template, JOB_TOKENS, tables, DOCUMENT_LISTS))

    assert "Delivery Fee" in out
    assert "35.00" in out


# --------------------------------------------------------------------------
# XML safety. A .docx is XML and values are spliced into it verbatim.
# --------------------------------------------------------------------------


def test_xml_special_characters_survive_the_merge():
    """Unescaped `&` used to be silent DATA LOSS, not a rendering glitch.

    "Smith & Sons <Pty> Ltd" merged as "Smith  Sons " — the ampersand opened
    an entity that never closed, and the malformed run was discarded when the
    document was reparsed. No error was raised anywhere.
    """
    template = make_docx(["Client: {{ Client }}"])
    tokens = [{"name": "Client", "value": "Smith & Sons <Pty> Ltd"}]
    out = text_of(render_docx(template, tokens))

    assert "Client: Smith & Sons <Pty> Ltd" in out


def test_entities_from_drupal_are_not_double_escaped():
    """Values come from Drupal fields that have been through strip_tags().

    Those still carry HTML entities, so "&amp;" is as likely as a raw "&".
    Decoding before escaping normalises both to the same correct output;
    without the decode this would render the literal text "&amp;".
    """
    template = make_docx(["Client: {{ Client }}"])
    tokens = [{"name": "Client", "value": "Smith &amp; Sons"}]
    out = text_of(render_docx(template, tokens))

    assert "Client: Smith & Sons" in out


def test_special_characters_survive_inside_a_list():
    """Escaping must not eat the separator the paragraph split depends on."""
    template = make_docx(["{{ Documents.Affidavit_Template }}"])
    lists = {"Documents.Affidavit_Template": ["Smith & Sons", "Jones <Pty> Ltd"]}
    out = paragraphs_of(render_docx(template, JOB_TOKENS, None, lists))

    assert out == ["Smith & Sons", "Jones <Pty> Ltd"]


# --------------------------------------------------------------------------
# line(): several fields on one line, without a gap where a value is missing.
# --------------------------------------------------------------------------


# One attendance note per row, as `resolveTableRowsForJob()` sends them. The
# middle two are the reason line() exists: real Jobs have notes with no
# employee recorded, and employees who attended without leaving a note.
ATTENDANCE_ROWS = {
    "Attendance_Notes": [
        {"Employee_Full_Name": "Example Officer", "Notes": "No response at the door."},
        {"Employee_Full_Name": "", "Notes": "Left a card."},
        {"Employee_Full_Name": "Second Officer", "Notes": ""},
        {"Employee_Full_Name": "Third Officer", "Notes": "Served in person."},
    ]
}

# Exactly what the token picker generates for Employee then Notes. Kept
# verbatim so a change to the JS generator that this service cannot render
# fails here rather than in Word.
ATTENDANCE_BLOCK = [
    ("{%p for row in Attendance_Notes %}", "Normal"),
    ("{{ line(row.Employee_Full_Name, row.Notes) }}", "List Number"),
    ("{%p endfor %}", "Normal"),
]


def test_line_pairs_each_row_on_its_own_line():
    out = paragraphs_of(render_docx(make_docx_styled(ATTENDANCE_BLOCK), [], ATTENDANCE_ROWS))

    assert out[0] == "Example Officer - No response at the door."
    assert out[3] == "Third Officer - Served in person."


def test_a_missing_value_takes_the_separator_with_it():
    """The bug this fixes: " - Left a card." and "Second Officer - "."""
    out = paragraphs_of(render_docx(make_docx_styled(ATTENDANCE_BLOCK), [], ATTENDANCE_ROWS))

    assert out[1] == "Left a card."
    assert out[2] == "Second Officer"
    assert not any(p.startswith(" - ") or p.endswith(" - ") for p in out)


def test_line_keeps_the_paragraph_style_so_numbering_survives():
    document = Document(io.BytesIO(render_docx(make_docx_styled(ATTENDANCE_BLOCK), [], ATTENDANCE_ROWS)))

    assert [p.style.name for p in document.paragraphs] == ["List Number"] * 4


def test_line_drops_a_whole_empty_row_to_nothing():
    """Neither value present: an empty line, not a bare separator."""
    rows = {"Attendance_Notes": [{"Employee_Full_Name": "", "Notes": ""}]}
    out = paragraphs_of(render_docx(make_docx_styled(ATTENDANCE_BLOCK), [], rows))

    assert out == [""]


def test_line_ignores_a_mistyped_token_rather_than_failing():
    """ChainableUndefined stringifies to "", so one bad name costs one slot."""
    template = make_docx(["{{ line(Client, Nonexistent_Field) }}"])
    out = paragraphs_of(render_docx(template, JOB_TOKENS))

    assert out == ["Example Statutory Authority"]


def test_line_works_outside_a_loop():
    """The no-repeating-parent case the picker also generates."""
    template = make_docx(["{{ line(Client, Client.Mail) }}"])
    out = paragraphs_of(render_docx(template, JOB_TOKENS))

    assert out == ["Example Statutory Authority - notices@example.test"]


def test_line_is_not_reported_as_an_undeclared_variable():
    """It is a global, so a template using it must not look like a bad token."""
    from render import find_placeholders

    found = find_placeholders(make_docx(["{{ line(Client, Client.Mail) }}"]))

    assert "line" not in found
    assert "Client" in found
